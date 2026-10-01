"""
generate_summary_report.py
Generates RENALDO_Update_Summary.docx — a plain-English summary of recent
work on the RENALDO dashboard, written for a non-technical audience.
Run with: python scripts/generate_summary_report.py
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Colours (matches the dashboard's accent palette) ──
NAVY      = RGBColor(0x10, 0x23, 0x3B)
ACCENT    = RGBColor(0x1D, 0x6F, 0xA5)
ACCENT_LT = RGBColor(0xE8, 0xF1, 0xF7)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
TEXT      = RGBColor(0x11, 0x18, 0x27)
TEXT_2    = RGBColor(0x52, 0x60, 0x6D)
TEXT_3    = RGBColor(0x8A, 0x94, 0xA3)


def para_shade(para, rgb_hex):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), rgb_hex)
    pPr.append(shd)


def add_para_border_left(para, rgb_hex="1D6FA5", width=18):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(width))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), rgb_hex)
    pBdr.append(left)
    pPr.append(pBdr)


doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.font.color.rgb = TEXT_2

for lvl, size, colour in [("Heading 1", 17, NAVY), ("Heading 2", 13, ACCENT)]:
    s = doc.styles[lvl]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = colour
    s.paragraph_format.space_before = Pt(16)
    s.paragraph_format.space_after = Pt(6)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def body(text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(8)
    p.add_run(text)
    return p


def bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        r.font.size = Pt(11)
        r.font.color.rgb = TEXT_2
    doc.add_paragraph()


def callout(text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    para_shade(p, "E8F1F7")
    add_para_border_left(p, "1D6FA5", 18)
    r = p.add_run(text)
    r.font.color.rgb = NAVY
    r.font.size = Pt(10.5)
    return p


# ════════════════════════════════════════════
# COVER
# ════════════════════════════════════════════
title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
para_shade(title_para, "10233B")
r = title_para.add_run("RENALDO")
r.font.name = "Calibri"
r.font.size = Pt(30)
r.font.bold = True
r.font.color.rgb = WHITE

sub_para = doc.add_paragraph()
sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
para_shade(sub_para, "10233B")
r = sub_para.add_run("Recent Updates — A Plain-English Summary")
r.font.name = "Calibri"
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0xB8, 0xD4, 0xE8)

org_para = doc.add_paragraph()
org_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
para_shade(org_para, "10233B")
r = org_para.add_run("UK Kidney Association  ·  RaDaR Registry  ·  September 2026")
r.font.name = "Calibri"
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0xB8, 0xD4, 0xE8)

doc.add_paragraph()

# ════════════════════════════════════════════
# INTRO
# ════════════════════════════════════════════
h1("What is this document?")
body(
    "RENALDO is the dashboard that shows how complete the RaDaR data is across all of the "
    "rare kidney disease patient groups. This document is a plain-English summary of the "
    "work carried out on it recently — what changed, why, and what's still being discussed. "
    "It avoids technical detail and is written for anyone, regardless of their background "
    "with the underlying software."
)

# ════════════════════════════════════════════
# 1. DATA CHANGES
# ════════════════════════════════════════════
h1("1. Changes agreed with the team")

h2("Follow-up measures")
bullets([
    "“Overall follow-up” now clearly shows how long each patient has been followed — "
    "from their diagnosis (or from when they joined RaDaR, if no diagnosis date is recorded) "
    "up to today, or to their date of death. Patients who have withdrawn from the study are "
    "no longer counted in this figure.",
    "A new measure, “Follow-up pre-KRT,” has been added. It shows how long patients were "
    "followed before starting kidney replacement therapy (dialysis or a transplant) — or up "
    "to today, for patients who haven't needed it yet. This appears for every condition group "
    "except CMV Post-Transplant and BK Nephropathy, where it isn't a meaningful measure.",
])

h2("Biochemistry (blood and urine test) information")
bullets([
    "The Biochemistry section has been removed from the CMV Post-Transplant and BK "
    "Nephropathy groups, as it doesn't apply to those patients.",
    "A new measure has been added showing how long, on average, patients had these tests "
    "recorded before starting kidney replacement therapy — giving a clearer picture of "
    "monitoring in the run-up to treatment.",
])

h2("Removing an internal-only group")
body(
    "The “Data Completeness” group has been removed from the public dashboard. It was "
    "only ever used for internal quality-checking and isn't a real group of patients, so it "
    "had no place on a page meant for external viewers such as researchers and collaborators."
)

# ════════════════════════════════════════════
# 2. LOOK AND FEEL
# ════════════════════════════════════════════
h1("2. Making the dashboard look and feel more professional")
bullets([
    "The whole visual design was refreshed — bright colours, gradients, and other decorative "
    "effects were replaced with a calmer, cleaner style more in keeping with a serious "
    "clinical and research tool.",
    "Fixed a display problem where longer descriptions in the tables were being cut off and "
    "hidden behind a scrollbar. Text now wraps onto a new line instead, so nothing is hidden.",
    "Fixed the “last refreshed” time shown at the top of the page, which was displaying an "
    "hour behind due to a technical time-zone issue.",
    "Fixed a colour mix-up where clicking on a table cell turned it red, which could be "
    "mistaken for a data problem — it now shows a neutral blue instead.",
    "Changed “Gender” to “Sex” for the relevant demographic field, using the plainer, "
    "less contentious clinical term.",
])

# ════════════════════════════════════════════
# 3. SECURITY / RELIABILITY
# ════════════════════════════════════════════
h1("3. Strengthening security and reliability behind the scenes")
body(
    "None of the following is visible on screen, but it makes the dashboard safer to run and "
    "easier to maintain going forward."
)
bullets([
    "Removed old, unused pieces of the program left over from earlier versions, reducing the "
    "chance of outdated logic causing confusion in future.",
    "Reorganised the program's inner workings into smaller, clearly-labelled pieces, and added "
    "over 50 automatic checks that catch mistakes before they can reach the live dashboard.",
    "Updated a security-sensitive software component that had known weaknesses to a safer, "
    "current version — tested directly against the live database to confirm nothing broke.",
    "Added basic protective measures to the website itself, reducing the risk of certain "
    "common web-based attacks.",
    "Set up an automatic check that runs every time the program is changed, to catch problems "
    "immediately rather than after they've gone live.",
])

# ════════════════════════════════════════════
# 4. STILL BEING DISCUSSED
# ════════════════════════════════════════════
h1("4. Still being discussed — not yet changed")
bullets([
    "Adding a short glossary box explaining terms like “Kidney Failure” and “Follow-up” "
    "once, instead of repeating the explanation in every section.",
    "Simplifying technical phrases (such as “test and control patients excluded”) so they "
    "make sense to people without a data background.",
    "Showing both “% complete” and “% missing” together at the top of each section, so no "
    "mental arithmetic is needed to compare figures.",
    "Deciding whether future changes to the live dashboard should be reviewed by a second "
    "person before going live.",
])

callout(
    "In short: the underlying data now reflects the team's latest decisions, the dashboard "
    "looks and behaves more professionally, and it's safer and easier to maintain — with a "
    "clear, agreed list of what's still open for discussion."
)

# ════════════════════════════════════════════
# FOOTER
# ════════════════════════════════════════════
doc.add_paragraph()
p = doc.add_paragraph(style="Normal")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("RENALDO Dashboard  ·  UK Kidney Association  ·  renaldo.onrender.com")
r.font.size = Pt(8.5)
r.font.color.rgb = TEXT_3

out = "RENALDO_Update_Summary.docx"
doc.save(out)
print(f"Saved: {out}")
